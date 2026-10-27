"""Core logic: fit a guest's name onto the invitation, and read/write guest lists."""

import os
import re
import sys
from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path
from typing import Optional

import openpyxl
from openpyxl.styles import Font
from PIL import Image, ImageDraw, ImageFont

# Works both from source and inside a PyInstaller bundle
BASE_DIR = Path(getattr(sys, "_MEIPASS", Path(__file__).resolve().parent))
FONTS_DIR = BASE_DIR / "fonts"
WINDOWS_FONTS_DIR = Path(os.environ.get("WINDIR", r"C:\Windows")) / "Fonts"

MIN_FONT_SIZE = 10
CHINESE_CHARS = re.compile(r"[\u3400-\u9fff\uf900-\ufaff]")


@dataclass(frozen=True)
class FontChoice:
    label: str
    path: Path
    weight: Optional[int] = None     # for variable fonts whose default weight is too thin
    chinese: Optional[bool] = False  # None = unknown (user's own font)


BUILT_IN_FONTS = [
    FontChoice("Alex Brush  (script)", FONTS_DIR / "AlexBrush-Regular.ttf"),
    FontChoice("Great Vibes  (script)", FONTS_DIR / "GreatVibes-Regular.ttf"),
    FontChoice("Parisienne  (script)", FONTS_DIR / "Parisienne-Regular.ttf"),
    FontChoice("Pinyon Script  (script)", FONTS_DIR / "PinyonScript-Regular.ttf"),
    FontChoice("Allura  (script)", FONTS_DIR / "Allura-Regular.ttf"),
    FontChoice("Dancing Script  (casual script)", FONTS_DIR / "DancingScript-Variable.ttf"),
    FontChoice("Cormorant Garamond  (classic)", FONTS_DIR / "CormorantGaramond-Variable.ttf", weight=500),
    FontChoice("Playfair Display  (classic)", FONTS_DIR / "PlayfairDisplay-Variable.ttf"),
    FontChoice("Lato  (modern)", FONTS_DIR / "Lato-Regular.ttf"),
]

# Chinese fonts are too large to bundle, so use the ones Windows ships with
SYSTEM_CHINESE_FONTS = [
    FontChoice("KaiTi 楷体  (Chinese brush)", WINDOWS_FONTS_DIR / "simkai.ttf", chinese=True),
    FontChoice("Microsoft YaHei 微软雅黑  (Chinese modern)", WINDOWS_FONTS_DIR / "msyh.ttc", chinese=True),
]


def available_fonts():
    return BUILT_IN_FONTS + [f for f in SYSTEM_CHINESE_FONTS if f.path.exists()]


def chinese_fallback():
    return next((f for f in SYSTEM_CHINESE_FONTS if f.path.exists()), None)


def has_chinese(text):
    return bool(CHINESE_CHARS.search(text))


@dataclass
class Placement:
    """Where the name goes, in pixels of the original image."""

    x_center: int = 850
    base_y: int = 2440
    max_y: int = 2470
    x_min: int = 230
    x_max: int = 1450
    max_font_size: int = 80
    color: str = "#000000"


@lru_cache(maxsize=128)
def load_font(choice, size):
    font = ImageFont.truetype(str(choice.path), size)
    if choice.weight:
        font.set_variation_by_axes([choice.weight])
    return font


def fit_font(draw, name, choice, p):
    """Shrink the font until the centred name fits between x_min and x_max."""
    font_size = p.max_font_size
    while True:
        font = load_font(choice, font_size)
        bbox = draw.textbbox((0, 0), name, font=font)
        text_width = bbox[2] - bbox[0]
        start_x = p.x_center - (text_width // 2)
        if (start_x >= p.x_min and start_x + text_width <= p.x_max) or font_size <= MIN_FONT_SIZE:
            return font, font_size, start_x
        font_size -= 2


def render_card(base_image, name, choice, p):
    # Latin-only fonts can't draw Chinese characters, so switch those names to a Chinese font
    if choice.chinese is False and has_chinese(name):
        choice = chinese_fallback() or choice

    img = base_image.copy()
    draw = ImageDraw.Draw(img)
    font, font_size, start_x = fit_font(draw, name, choice, p)

    # Nudge the text down slightly as the font gets smaller
    reduction_steps = (p.max_font_size - font_size) // 10
    adjusted_y = min(p.base_y + reduction_steps * 20, p.max_y)

    draw.text((start_x, adjusted_y), name, fill=p.color, font=font)
    return img


def read_names(path):
    """Names from column A of the first sheet; a 'Name' header row is skipped."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        names = [str(row[0]).strip() for row in wb.active.iter_rows(max_col=1, values_only=True) if row and row[0] is not None]
    finally:
        wb.close()
    if names and names[0].lower() == "name":
        names = names[1:]
    return [n for n in names if n]


def write_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Guests"
    ws["A1"] = "Name"
    ws["A1"].font = Font(bold=True)
    ws["C1"] = "One guest (or couple) per row in column A. Replace the examples below."
    ws["C1"].font = Font(italic=True, color="808080")
    for row, example in enumerate(["John Doe", "Mary Tan and David Lee", "Uncle Robert"], start=2):
        ws.cell(row=row, column=1, value=example)
    ws.column_dimensions["A"].width = 40
    wb.save(path)


def safe_filename(name, prefix="Invite_"):
    cleaned = re.sub(r'[\\/:*?"<>|]', "", name).strip()
    return f"{prefix}{cleaned}.jpg"
